import io, asyncio
import openpyxl
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
from googleapiclient.errors import HttpError
from config.settings import EXCEL_FILE_ID, SHEET_NAME, get_ist_now
from config.logger import log_error, logger

class DriveAccessError(Exception): pass

SCOPES = ['https://www.googleapis.com/auth/drive']
creds = Credentials.from_service_account_file('credentials.json', scopes=SCOPES)
drive_service = build('drive', 'v3', credentials=creds, cache_discovery=False)

def _pull():
    try:
        req = drive_service.files().get_media(fileId=EXCEL_FILE_ID)
        buf = io.BytesIO()
        downloader = MediaIoBaseDownload(buf, req)
        done = False
        while not done: _, done = downloader.next_chunk()
        buf.seek(0)
        return openpyxl.load_workbook(buf)
    except HttpError as e:
        if e.resp.status == 404:
            logger.error(f"[ERROR] DRIVE 404: Cannot find file {EXCEL_FILE_ID}. Did you share it with the Service Account email?")
            raise DriveAccessError("I cannot access the Google Drive file. Please ensure it is shared with my Service Account email as Editor.")
        raise e

def _push(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    media = MediaIoBaseUpload(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', resumable=True)
    drive_service.files().update(fileId=EXCEL_FILE_ID, media_body=media).execute()

async def get_categories():
    def fetch():
        try:
            wb = _pull(); ws = wb[SHEET_NAME]
            return [str(ws[f"A{r}"].value).strip().title() for r in range(12, 17) if ws[f"A{r}"].value]
        except DriveAccessError as e:
            raise e
        except Exception as e:
            log_error("Failed fetching categories", e)
            return []
    return await asyncio.to_thread(fetch)

async def get_data():
    def fetch():
        try:
            wb = _pull(); ws = wb[SHEET_NAME]
            data, row = [], 18
            while ws[f"C{row}"].value:
                data.append({
                    "rowId": row, 
                    "expenseDate": str(ws[f"C{row}"].value), 
                    "itemName": str(ws[f"D{row}"].value).title(), 
                    "categoryName": str(ws[f"E{row}"].value).title(), 
                    "expenseAmount": float(ws[f"F{row}"].value or 0)
                })
                row += 1
            return data
        except DriveAccessError as e:
            raise e
        except Exception as e:
            log_error("Failed fetching data", e)
            return []
    return await asyncio.to_thread(fetch)

async def save_expense(item, cat, amt, date=None):
    def save():
        try:
            wb = _pull(); ws = wb[SHEET_NAME]
            row = 18
            while ws[f"C{row}"].value: row += 1
            ws[f"C{row}"] = date or get_ist_now().strftime("%d-%m-%Y")
            ws[f"D{row}"] = item.title(); ws[f"E{row}"] = cat.title(); ws[f"F{row}"] = float(amt); ws[f"G{row}"] = "ACTIVE"
            _push(wb)
        except DriveAccessError as e:
            raise e
        except Exception as e:
            log_error(f"Failed inserting expense: {item}", e)
            raise e
    await asyncio.to_thread(save)

async def hard_delete(rows):
    def delete():
        try:
            wb = _pull(); ws = wb[SHEET_NAME]
            for r in sorted(rows, reverse=True): ws.delete_rows(r, 1)
            _push(wb)
        except DriveAccessError as e:
            raise e
        except Exception as e:
            log_error("Failed hard deleting rows", e)
            raise e
    await asyncio.to_thread(delete)
